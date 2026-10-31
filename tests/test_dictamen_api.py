"""API de dictámenes técnicos, límites por ruta, store_file, despliegue y lógica del frontend.

El extractor (Qwen) siempre es un doble: estas pruebas no llaman al LLM. Usan las piezas
reales de app/dictamen (schema, tecnicos, extract, generate y la plantilla); el flujo completo
con OficioExtractor real y el LLM simulado está en test_dictamen_integracion.py.
"""

from __future__ import annotations

import importlib
import io
import json
import os
import re
import shutil
import subprocess
import time
from pathlib import Path

import httpx
import pytest
import yaml
from fastapi.testclient import TestClient
from starlette.responses import JSONResponse

from app.documents.connector import DocumentServiceError
from app.limits import TOO_LARGE, RequestLimits
from app.main import app
from app.prompts import SYSTEM_PROMPT
from app.settings import Settings

ROOT = Path(__file__).resolve().parent.parent
APP_JS = ROOT / "app" / "static" / "app.js"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

OFICIO_DATA = {
    "oficio_recibido": "FGEO/DA/0532/2026",
    "fecha_oficio_recibido": "2026-08-20",
    "cargo_remitente": "Directora Administrativa",
    "nombre_remitente": "Mtra. Laura Isabel Mendoza Cruz",
    "equipos": [
        {"tipo_equipo": "CPU", "marca": "HP", "modelo": "ProDesk 400", "numero_serie": "MXL1234ABC",
         "numero_inventario": "FGEO-CPU-004512", "motivo": None},
        {"tipo_equipo": "Monitor", "marca": None, "modelo": "E2216H", "numero_serie": "CN0ABC123",
         "numero_inventario": None, "motivo": None},
    ],
}

FULL = {
    "data": {
        "oficio_recibido": "FGEO/DA/0532/2026",
        "fecha_oficio_recibido": "2026-08-20",
        "cargo_remitente": "Directora Administrativa",
        "nombre_remitente": "Mtra. Laura Isabel Mendoza Cruz",
        "equipos": [
            {"tipo_equipo": "CPU", "marca": "HP", "modelo": "ProDesk 400", "numero_serie": "MXL1234ABC",
             "numero_inventario": "FGEO-CPU-004512", "motivo": None},
            {"tipo_equipo": "Monitor", "marca": "Dell", "modelo": "E2216H", "numero_serie": "CN0ABC123",
             "numero_inventario": None, "motivo": "Pantalla rota."},
        ],
    },
    "own": {
        "numero_oficio": "FGEO/CGSI/1429/2026",
        "dictamen_inicial": "191",
        "iniciales": "MDTG/IGF, dsa",
        "motivo": "Equipo con falla de encendido.",
    },
}


# ---------------------------------------------------------------- dobles y utilidades
class StubExtractor:
    def __init__(self, result=None, error=None):
        self.result = result
        self.error = error
        self.calls = []

    async def extract(self, content: bytes, filename: str):
        self.calls.append((content, filename))
        if self.error:
            raise self.error
        return self.result


def tiny_pdf() -> bytes:
    import pymupdf

    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "OFICIO NUM. FGEO/DA/0532/2026")
    return doc.tobytes()


def _limits_of(asgi_app) -> RequestLimits | None:
    node = asgi_app.middleware_stack
    while node is not None and not isinstance(node, RequestLimits):
        node = getattr(node, "app", None)
    return node


def _requires(*modules):
    # Las piezas ya están integradas: si alguna no importa, la prueba debe fallar, no omitirse.
    return [importlib.import_module(m) for m in modules]


@pytest.fixture
def client(connector):
    """App real con el conector simulado; los dictámenes activos aunque falte alguna pieza."""
    with TestClient(app) as c:
        limits = _limits_of(c.app)
        if limits is not None:
            limits.recent.clear()  # la suite comparte el presupuesto de 30 POST/min de la app
        c.app.state.connector = connector
        c.app.state.dictamen_enabled = True
        c.app.state.dictamen_extractor = StubExtractor({"data": OFICIO_DATA, "mode": "texto", "pages": 1, "elapsed_s": 6.5})
        original = c.app.state.settings
        yield c
        c.app.state.settings = original


def _settings(c, **update):
    c.app.state.settings = c.app.state.settings.model_copy(update=update)


# ---------------------------------------------------------------- límites por ruta
async def _echo(scope, receive, send):
    body = await receive()
    await JSONResponse({"size": len(body["body"])})(scope, receive, send)


async def test_upload_route_has_its_own_cap_and_other_posts_keep_theirs():
    guarded = RequestLimits(_echo, max_bytes=10, route_limits={"/api/dictamen/oficio": (100, "Archivo muy grande.")})
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=guarded), base_url="http://test") as c:
        assert (await c.post("/api/dictamen/oficio", content=b"x" * 80)).json()["size"] == 80
        r = await c.post("/api/chat", content=b"x" * 80)
        assert r.status_code == 413 and r.json()["message"] == TOO_LARGE
        r = await c.post("/api/dictamen/oficio", content=b"x" * 150)
        assert r.status_code == 413 and r.json()["message"] == "Archivo muy grande."

        async def chunks():
            for _ in range(3):
                yield b"x" * 40
        r = await c.post("/api/dictamen/oficio", content=chunks())  # chunked, sin Content-Length
        assert r.status_code == 413


async def test_declared_oversize_is_rejected_without_reading_the_body():
    guarded = RequestLimits(_echo, max_bytes=10, route_limits={"/api/dictamen/oficio": (100, "Archivo muy grande.")})
    sent, reads = [], []

    async def receive():
        reads.append(1)
        return {"type": "http.request", "body": b"", "more_body": False}

    async def send(message):
        sent.append(message)

    scope = {"type": "http", "method": "POST", "path": "/api/dictamen/oficio",
             "headers": [(b"content-length", str(50 * 1024 * 1024).encode())]}
    await guarded(scope, receive, send)
    assert sent[0]["status"] == 413 and not reads
    assert guarded.active == 0  # la plaza se libera


def test_app_registers_upload_limit_from_settings(client):
    limits = _limits_of(client.app)
    s = client.app.state.settings
    cap, message = limits.route_limits["/api/dictamen/oficio"]
    assert cap > s.dictamen_max_upload_bytes and limits.max_bytes == s.max_request_bytes
    assert f"{s.dictamen_max_upload_mb:g} MB" in message


# ---------------------------------------------------------------- store_file
def test_store_file_is_downloadable_and_respects_retention(connector):
    doc = connector.store_file("dictamen_tecnico", "Dictamen técnico", b"PK\x03\x04docx", DOCX_MIME,
                               "../dictamen FGEO/CGSI 1.docx")
    assert doc.filename == "CGSI 1.docx"  # sin rutas
    assert doc.type == "dictamen_tecnico" and doc.mime_type == DOCX_MIME
    assert doc.download_path == f"/api/documents/{doc.id}"
    path, meta = connector.get(doc.id)
    assert path.read_bytes() == b"PK\x03\x04docx" and meta.title == "Dictamen técnico"
    old = time.time() - 25 * 3600
    for p in connector.storage_dir.glob(f"{doc.id}.*"):
        os.utime(p, (old, old))
    assert connector.get(doc.id) is None
    with pytest.raises(DocumentServiceError):
        connector.store_file("dictamen_tecnico", "Dictamen técnico", b"", DOCX_MIME, "vacio.docx")


# ---------------------------------------------------------------- configuración y despliegue
def test_settings_defaults_and_model_fallback(monkeypatch):
    for var in ("DICTAMEN_ENABLED", "DICTAMEN_EXTRACT_MODEL", "DICTAMEN_TEMPLATE_PATH", "DICTAMEN_MAX_UPLOAD_MB"):
        monkeypatch.delenv(var, raising=False)
    s = Settings(_env_file=None)
    assert s.dictamen_enabled is True and s.dictamen_max_upload_mb == 15 and s.dictamen_max_pages == 4
    assert s.dictamen_extract_timeout_s == 240 and s.dictamen_model == s.llm_model
    assert s.dictamen_template_path == ROOT / "config" / "dictamen" / "dictamen.docx"
    assert s.dictamen_tecnicos_path == ROOT / "config" / "dictamen" / "tecnicos.json"
    monkeypatch.setenv("DICTAMEN_EXTRACT_MODEL", "")
    monkeypatch.setenv("DICTAMEN_TEMPLATE_PATH", "otra/plantilla.docx")
    monkeypatch.setenv("DICTAMEN_ENABLED", "false")
    s = Settings(_env_file=None)
    assert s.dictamen_extract_model is None and s.dictamen_template_path == ROOT / "otra" / "plantilla.docx"
    assert s.dictamen_enabled is False
    monkeypatch.setenv("DICTAMEN_EXTRACT_MODEL", "qwen-vl")
    assert Settings(_env_file=None).dictamen_model == "qwen-vl"


def test_public_demo_compose_disables_dictamen_by_default():
    compose = yaml.safe_load((ROOT / "compose.yaml").read_text())
    env = compose["services"]["app"]["environment"]
    assert env["DEMO_MODE"] == "true"
    assert env["DICTAMEN_ENABLED"] == "${DICTAMEN_ENABLED:-false}"


def test_system_prompt_points_to_the_attach_button():
    assert "📎" in SYSTEM_PROMPT and "dictámenes técnicos" in SYSTEM_PROMPT.lower()
    SYSTEM_PROMPT.format(today="hoy", documents="ninguno")  # sigue siendo formateable


# ---------------------------------------------------------------- endpoints: configuración
def test_config_exposes_fields_tecnicos_and_limits(client):
    _requires("app.dictamen.schema", "app.dictamen.tecnicos")
    cfg = client.get("/api/dictamen/config").json()
    assert cfg["enabled"] is True and cfg["accept"] == ".pdf,.docx,.jpg,.jpeg,.png"
    assert cfg["max_upload_mb"] == client.app.state.settings.dictamen_max_upload_mb
    assert [f["name"] for f in cfg["fields"]["general"]] == [
        "oficio_recibido", "fecha_oficio_recibido", "cargo_remitente", "nombre_remitente"]
    assert {"tipo_equipo", "numero_inventario", "motivo"} <= {f["name"] for f in cfg["fields"]["equipo"]}
    assert "igf" in [t["iniciales"] for t in cfg["tecnicos"]] and cfg["max_equipos"] == 100
    assert client.get("/api/config").json()["dictamen"] == cfg


def test_disabled_dictamen_hides_everything(client):
    client.app.state.dictamen_enabled = False
    assert client.get("/api/config").json()["dictamen"] is None
    assert client.get("/api/dictamen/config").status_code == 404
    r = client.post("/api/dictamen/oficio", files={"file": ("oficio.pdf", b"%PDF-1.4", "application/pdf")})
    assert r.status_code == 404
    assert client.post("/api/dictamen/generar", json=FULL).status_code == 404


def test_dictamen_not_loaded_when_setting_is_false():
    import app.main as main

    assert main._load_dictamen(Settings(dictamen_enabled=False)) is False


# ---------------------------------------------------------------- endpoints: oficio
def test_upload_returns_all_extracted_data_and_faltantes(client):
    _requires("app.dictamen.schema", "app.dictamen.extract")
    pdf = tiny_pdf()
    r = client.post("/api/dictamen/oficio", files={"file": ("../../Oficio <b>.pdf", pdf, "text/plain")},
                    data={"conversation_id": "abc"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["filename"] == "Oficio <b>.pdf"  # sin rutas; el frontend lo escapa
    assert body["mode"] == "texto" and body["pages"] == 1 and body["total_pages"] == 1 and body["elapsed_s"] == 6.5
    assert body["data"]["fecha_oficio_recibido"] == "20 DE AGOSTO DEL 2026"  # ISO normalizada
    assert len(body["data"]["equipos"]) == 2
    assert set(body["data"]["equipos"][0]) == {"tipo_equipo", "marca", "modelo", "numero_serie", "numero_inventario", "motivo"}
    assert body["own"] == {"numero_oficio": None, "dictamen_inicial": None, "iniciales": None, "motivo": None}
    assert body["faltantes"] == ["equipo_2_marca", "numero_oficio", "dictamen_inicial", "iniciales", "motivo"]
    content, filename = client.app.state.dictamen_extractor.calls[0]
    assert content == pdf and filename == "Oficio <b>.pdf"


def test_upload_rejects_unknown_formats_empty_and_oversized_files(client):
    _requires("app.dictamen.schema", "app.dictamen.extract")
    stub = client.app.state.dictamen_extractor
    r = client.post("/api/dictamen/oficio", files={"file": ("oficio.pdf", b"no soy un pdf", "application/pdf")})
    assert r.status_code == 415 and "PDF" in r.json()["message"]
    assert client.post("/api/dictamen/oficio", files={"file": ("vacio.pdf", b"", "application/pdf")}).status_code == 415
    _settings(client, dictamen_max_upload_mb=0.001)  # ~1 KB
    r = client.post("/api/dictamen/oficio", files={"file": ("grande.pdf", b"%PDF-1.4" + b"0" * 4096, "application/pdf")})
    assert r.status_code == 413 and "MB" in r.json()["message"]
    assert stub.calls == []  # nunca llegó al LLM


def test_upload_extraction_error_returns_empty_editable_draft(client):
    schema, extract = _requires("app.dictamen.schema", "app.dictamen.extract")
    client.app.state.dictamen_extractor = StubExtractor(error=extract.ExtractionError("El modelo no respondió a tiempo."))
    r = client.post("/api/dictamen/oficio", files={"file": ("oficio.pdf", tiny_pdf(), "application/pdf")})
    assert r.status_code == 502
    body = r.json()
    draft = schema.empty_draft()
    assert body["message"] == "El modelo no respondió a tiempo."
    assert body["data"] == draft["data"] and body["own"] == draft["own"]
    assert "equipos" in body["faltantes"] and body["filename"] == "oficio.pdf"
    client.app.state.dictamen_extractor = StubExtractor(error=RuntimeError("detalle interno /srv/x"))
    r = client.post("/api/dictamen/oficio", files={"file": ("oficio.pdf", tiny_pdf(), "application/pdf")})
    assert r.status_code == 502 and "/srv" not in r.json()["message"]


# ---------------------------------------------------------------- endpoints: generar
def test_generar_rejects_wrong_types_long_strings_and_too_many_equipos(client):
    schema, *_ = _requires("app.dictamen.schema", "app.dictamen.tecnicos", "app.dictamen.generate", "app.dictamen.extract")
    cases = [
        {**FULL, "data": {**FULL["data"], "equipos": "3"}},
        {**FULL, "data": {**FULL["data"], "nombre_remitente": "x" * 501}},
        {**FULL, "own": {**FULL["own"], "motivo": {"a": 1}}},
        {**FULL, "data": {**FULL["data"], "equipos": [FULL["data"]["equipos"][0]] * (schema.MAX_EQUIPOS + 1)}},
        {**FULL, "data": {**FULL["data"], "equipos": ["texto"]}},
    ]
    for payload in cases:
        r = client.post("/api/dictamen/generar", json=payload)
        assert r.status_code == 422, payload
        assert r.json()["message"] and r.json()["errors"]


def test_generar_reports_faltantes_with_bot_names(client):
    _requires("app.dictamen.schema", "app.dictamen.tecnicos", "app.dictamen.generate", "app.dictamen.extract")
    r = client.post("/api/dictamen/generar", json={**FULL, "data": {**FULL["data"], "equipos": []}})
    assert r.status_code == 422 and r.json()["faltantes"] == ["equipos"]
    assert "al menos un equipo" in r.json()["message"]
    r = client.post("/api/dictamen/generar", json={**FULL, "own": {**FULL["own"], "dictamen_inicial": "19a", "numero_oficio": " "}})
    assert r.status_code == 422 and r.json()["faltantes"] == ["numero_oficio", "dictamen_inicial"]
    # motivo general obligatorio solo si algún equipo no tiene el suyo
    eqs = [{**e, "motivo": "Falla"} for e in FULL["data"]["equipos"]]
    r = client.post("/api/dictamen/generar", json={"data": {**FULL["data"], "equipos": eqs}, "own": {**FULL["own"], "motivo": ""}})
    assert r.status_code != 422 or "motivo" not in r.json().get("faltantes", [])


def test_generar_rejects_unknown_tecnico_when_catalog_has_names(client, tmp_path):
    _requires("app.dictamen.schema", "app.dictamen.tecnicos", "app.dictamen.generate", "app.dictamen.extract")
    catalog = tmp_path / "tecnicos.json"
    catalog.write_text(json.dumps([{"nombre": "Ivan García Flores", "iniciales": "igf"}]), encoding="utf-8")
    _settings(client, dictamen_tecnicos_path=catalog)
    r = client.post("/api/dictamen/generar", json={**FULL, "own": {**FULL["own"], "iniciales": "zzz"}})
    assert r.status_code == 422 and r.json()["field"] == "iniciales"
    assert "igf" in r.json()["message"] and r.json()["errors"]


def test_generar_stores_document_with_stubbed_render(client, monkeypatch, tmp_path):
    _, _, generate, extract = _requires("app.dictamen.schema", "app.dictamen.tecnicos", "app.dictamen.generate", "app.dictamen.extract")
    template = tmp_path / "plantilla.docx"
    template.write_bytes(b"PK")
    _settings(client, dictamen_template_path=template)
    seen = {}

    def fake_render(data, own, template_path, today=None):
        seen.update(data=data, own=own, template=template_path, today=today)
        return b"PK\x03\x04fake-docx"

    monkeypatch.setattr(generate, "render_dictamen", fake_render)
    r = client.post("/api/dictamen/generar", json={**FULL, "conversation_id": "c1"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["numeros"] == [191, 192] and body["iniciales"] == "igf/dsa"  # sin MDTG, normalizadas
    assert re.fullmatch(r"dictamen_FGEO_CGSI_1429_2026_\d{4}-\d{2}-\d{2}\.docx", body["filename"])
    assert body["mime_type"] == extract.DOCX_MIME == DOCX_MIME and body["title"] == "Dictamen técnico"
    assert seen["data"]["fecha_oficio_recibido"] == "20 DE AGOSTO DEL 2026" and seen["template"] == template
    assert set(seen["data"]["equipos"][1]) == {"tipo_equipo", "marca", "modelo", "numero_serie", "numero_inventario", "motivo"}
    assert body["filename"].endswith(f"{seen['today'].isoformat()}.docx")
    f = client.get(body["url"])
    assert f.status_code == 200 and f.content == b"PK\x03\x04fake-docx"
    assert f.headers["content-type"] == DOCX_MIME and f.headers["cache-control"] == "private, no-store"

    monkeypatch.setattr(generate, "render_dictamen", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("boom /srv")))
    r = client.post("/api/dictamen/generar", json=FULL)
    assert r.status_code == 500 and "/srv" not in r.json()["message"]


def test_generar_without_template_is_a_clear_server_error(client, tmp_path):
    _requires("app.dictamen.schema", "app.dictamen.tecnicos", "app.dictamen.generate", "app.dictamen.extract")
    _settings(client, dictamen_template_path=tmp_path / "no-existe.docx")
    r = client.post("/api/dictamen/generar", json=FULL)
    assert r.status_code == 503 and "plantilla" in r.json()["message"]


def test_generar_with_real_template_produces_complete_docx(client):
    _requires("app.dictamen.schema", "app.dictamen.tecnicos", "app.dictamen.generate", "app.dictamen.extract")
    assert client.app.state.settings.dictamen_template_path.exists(), "python scripts/convert_dictamen_template.py"
    from docx import Document

    payload = {**FULL, "own": {**FULL["own"], "dictamen_inicial": "0191"}}
    r = client.post("/api/dictamen/generar", json=payload)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["numeros"] == ["0191", "0192"]  # como se imprimen
    f = client.get(body["url"])
    assert f.headers["content-type"] == DOCX_MIME
    doc = Document(io.BytesIO(f.content))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    text = "\n".join(parts)
    for expected in ["FGEO/CGSI/1429/2026", "FGEO/DA/0532/2026", "20 DE AGOSTO DEL 2026", "MXL1234ABC",
                     "CN0ABC123", "S/N", "FGEO-CPU-004512", "PANTALLA ROTA", "MDTG/igf/dsa", "0191", "0192"]:
        assert expected in text, expected
    assert "PANTALLA ROTA.." not in text and "{{" not in text and "{%" not in text


# ---------------------------------------------------------------- frontend (node)
NODE = shutil.which("node")


def _pure_js() -> str:
    src = APP_JS.read_text(encoding="utf-8")
    start, end = src.index("// DICTAMEN-PURE-START"), src.index("// DICTAMEN-PURE-END")
    return src[start:end]


def _run_node(script: str) -> dict:
    out = subprocess.run([NODE, "-e", script], capture_output=True, text=True, timeout=30, check=True)
    return json.loads(out.stdout)


def test_frontend_parses():
    if not NODE:
        pytest.skip("node no está instalado")
    subprocess.run([NODE, "--check", str(APP_JS)], check=True, timeout=30)
    html = (ROOT / "app" / "static" / "index.html").read_text(encoding="utf-8")
    assert 'id="attach"' in html and 'id="attach-input"' in html and "hidden" in html.split('id="attach"')[1][:200]


FIELDS = {
    "general": [{"name": n, "label": n, "required": True} for n in
                ["oficio_recibido", "fecha_oficio_recibido", "cargo_remitente", "nombre_remitente"]],
    "equipo": [{"name": "tipo_equipo", "label": "Tipo", "required": True}, {"name": "marca", "label": "Marca", "required": True},
               {"name": "modelo", "label": "Modelo", "required": True}, {"name": "numero_serie", "label": "Serie", "required": True},
               {"name": "numero_inventario", "label": "Inventario", "required": False},
               {"name": "motivo", "label": "Motivo", "required": False}],
    "own": [{"name": "numero_oficio", "label": "No.", "required": True}, {"name": "dictamen_inicial", "label": "Inicial", "required": True},
            {"name": "iniciales", "label": "Iniciales", "required": True}, {"name": "motivo", "label": "Motivo", "required": True}],
}

FALTANTES_CASES = [
    {"data": {"equipos": []}, "own": {}},
    {"data": OFICIO_DATA, "own": {}},
    {"data": {**OFICIO_DATA, "equipos": [{**e, "motivo": "x"} for e in OFICIO_DATA["equipos"]]},
     "own": {"numero_oficio": "A", "dictamen_inicial": " 0191 ", "iniciales": "igf"}},
    {"data": FULL["data"], "own": {**FULL["own"], "dictamen_inicial": "19.5", "iniciales": "   "}},
    {"data": FULL["data"], "own": FULL["own"]},
    {"data": {**FULL["data"], "oficio_recibido": "  ", "equipos": [{"tipo_equipo": "CPU"}]}, "own": {"motivo": "m"}},
]


def test_frontend_card_logic_with_node():
    if not NODE:
        pytest.skip("node no está instalado")
    script = _pure_js() + r"""
const fields = %s;
const cases = %s;
const out = { faltantes: cases.map((c) => dictFaltantes(fields, c.data, c.own)) };
const cfg = { fields, tecnicos: [{ nombre: "", iniciales: "igf" }], max_equipos: 3 };
const c = dictNewCard({ filename: "o.pdf", data: { oficio_recibido: "X", equipos: [{ marca: null, extra: "no" }, "basura"] },
  own: { numero_oficio: "N" }, mode: "imagen", pages: 2, elapsed_s: 16.2 }, cfg, "card1");
out.initial = JSON.parse(JSON.stringify({ data: c.data, own: c.own, detected: c.detected, status: c.status }));
out.editEquipos = dictApplyEdit(c, "data", "equipos", 0, "3");   // lista blanca: no se puede
out.editUnknown = dictApplyEdit(c, "eq", "otros_componentes", 0, "x");
out.editRowMissing = dictApplyEdit(c, "eq", "marca", 5, "x");
c.invalid = ["equipo_1_marca", "iniciales"]; c.error = "viejo";
out.editOk = dictApplyEdit(c, "eq", "marca", 0, "  HP  ");
out.afterEdit = { invalid: c.invalid, error: c.error };
out.add = [dictAddEquipo(c), dictAddEquipo(c), dictAddEquipo(c)];   // máximo 3
dictApplyEdit(c, "eq", "numero_serie", 1, "SER-2");
out.remove = [dictRemoveEquipo(c, 7), dictRemoveEquipo(c, 2)];
dictApplyEdit(c, "own", "dictamen_inicial", null, "0099");
out.payload = dictPayload(c, null);
out.rango = [dictRango("0099", 2), dictRango("191", 1), dictRango("19a", 2), dictRango("191", 0), dictRango("9", 3)];
out.summary = [dictSummary(2, "a<b>.pdf"), dictSummary(0, "x.pdf")];
out.result = [dictResultText({ numeros: [191, 193], iniciales: "igf/dsa" }), dictResultText({ numeros: ["0191", "0191"], iniciales: "igf" })];
const entryReading = { upload: "reading", filename: "f.pdf", busyText: "Leyendo", text: "" };
const entrySending = { dictamen: { status: "sending" } };
const entryDone = { dictamen: { status: "done" } };
[entryReading, entrySending, entryDone].forEach(dictOnLoad);
out.onLoad = [entryReading, entrySending.dictamen.status, entryDone.dictamen.status];
c.status = "done"; c.result = { numeros: [99, 100], iniciales: "igf", filename: "d.docx", equipos: 2 };
out.noteDone = dictHistoryNote(c);
c.manual = true; c.status = "cancelled";
out.noteManual = dictHistoryNote(c);
out.errors = [dictErrorText({ message: "Revisa", errors: ["a", "Revisa"] }, 422), dictErrorText({ detail: "No encontrado" }, 404),
              dictErrorText(null, 500), dictErrorText({ detail: [{ msg: "x" }] }, 422)];
out.sentence = [dictSentence("Hola"), dictSentence("Listo."), dictSentence("")];
console.log(JSON.stringify(out));
""" % (json.dumps(FIELDS), json.dumps(FALTANTES_CASES))
    out = _run_node(script)

    assert out["faltantes"][0] == ["oficio_recibido", "fecha_oficio_recibido", "cargo_remitente", "nombre_remitente",
                                   "equipos", "numero_oficio", "dictamen_inicial", "iniciales"]
    assert out["faltantes"][1] == ["equipo_2_marca", "numero_oficio", "dictamen_inicial", "iniciales", "motivo"]
    assert out["faltantes"][2] == ["equipo_2_marca"]
    assert out["faltantes"][3] == ["dictamen_inicial", "iniciales"]
    assert out["faltantes"][4] == []
    assert out["faltantes"][5] == ["oficio_recibido", "equipo_1_marca", "equipo_1_modelo", "equipo_1_numero_serie",
                                   "numero_oficio", "dictamen_inicial", "iniciales"]

    init = out["initial"]
    assert init["status"] == "pending" and init["detected"] == 1
    assert init["data"]["oficio_recibido"] == "X" and init["data"]["cargo_remitente"] == ""
    assert init["data"]["equipos"] == [{f["name"]: "" for f in FIELDS["equipo"]}]  # sin claves ajenas
    assert init["own"] == {"numero_oficio": "N", "dictamen_inicial": "", "iniciales": "", "motivo": ""}
    assert out["editEquipos"] is False and out["editUnknown"] is False and out["editRowMissing"] is False
    assert out["editOk"] is True and out["afterEdit"] == {"invalid": ["iniciales"], "error": ""}
    assert out["add"] == [True, True, False]
    assert out["remove"] == [False, True]
    payload = out["payload"]
    assert set(payload) == {"data", "own", "conversation_id"} and payload["conversation_id"] is None
    assert payload["data"]["oficio_recibido"] == "X" and payload["data"]["cargo_remitente"] is None
    assert [e["marca"] for e in payload["data"]["equipos"]] == ["HP", None]
    assert payload["data"]["equipos"][1]["numero_serie"] == "SER-2"
    assert all(set(e) == {f["name"] for f in FIELDS["equipo"]} for e in payload["data"]["equipos"])
    assert payload["own"] == {"numero_oficio": "N", "dictamen_inicial": "0099", "iniciales": None, "motivo": None}
    assert out["rango"] == ["Se numerarán del 0099 al 0100.", "Se numerará como dictamen 191.", "", "", "Se numerarán del 9 al 11."]
    assert out["summary"][0] == "Detecté 2 equipo(s) en el oficio «a<b>.pdf». Revisa todos los datos y completa los marcados."
    assert out["summary"][1].startswith("No detecté equipos")
    assert out["result"] == ["Dictamen generado: dictámenes 191 a 193, iniciales MDTG/igf/dsa. Descárgalo abajo.",
                             "Dictamen generado: dictamen 0191, iniciales MDTG/igf. Descárgalo abajo."]
    reading, sending, done = out["onLoad"]
    assert reading["upload"] == "interrupted" and reading["busyText"] == "" and "f.pdf" in reading["text"]
    assert sending == "interrupted" and done == "done"
    assert out["noteDone"].startswith("(Sistema) Detecté 1 equipo(s) en el oficio «o.pdf»")
    assert "«d.docx»" in out["noteDone"] and "números 99 a 100" in out["noteDone"] and "MDTG/igf" in out["noteDone"]
    assert out["noteManual"].startswith("(Sistema) No se pudo leer el oficio «o.pdf»") and "descartó" in out["noteManual"]
    assert out["errors"] == ["Revisa · a", "No encontrado", "Error 500", "Error 422"]
    assert out["sentence"] == ["Hola.", "Listo.", ""]


def test_frontend_faltantes_match_python_schema():
    """La tarjeta recalcula los faltantes en el navegador con las mismas reglas que el servidor."""
    if not NODE:
        pytest.skip("node no está instalado")
    (schema,) = _requires("app.dictamen.schema")
    fields = schema.field_spec()
    script = _pure_js() + "\nconsole.log(JSON.stringify(%s.map((c) => dictFaltantes(%s, c.data, c.own))));" % (
        json.dumps(FALTANTES_CASES), json.dumps(fields))
    js = _run_node(script)
    assert js == [schema.faltantes(c["data"], c["own"]) for c in FALTANTES_CASES]


def test_frontend_shows_upload_messages_as_plain_text():
    """El nombre del archivo y los mensajes del servidor no pasan por el markdown del chat."""
    src = APP_JS.read_text(encoding="utf-8")
    assert "entry.plain ?" in src and "escapeHtml(entry.text)" in src
    upload = src[src.index("async function uploadOficio"):src.index("function renderAll")]
    assert "plain: true" in upload and "innerHTML" not in upload
