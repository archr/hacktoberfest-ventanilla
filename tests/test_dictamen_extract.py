import asyncio
import base64
import io
import json
import struct
import zipfile
import zlib
from pathlib import Path

import httpx
import pymupdf
import pytest
from PIL import Image

from app.dictamen import extract as E
from app.dictamen.extract import DOCX_MIME, ExtractionError, OficioExtractor, sniff_mime

FIX = Path(__file__).resolve().parent / "fixtures" / "dictamen"
TRUTH = json.loads((FIX / "oficio_ground_truth.json").read_text(encoding="utf-8"))
BASE = "http://llm.test/v1"


def fixture(name: str) -> bytes:
    return (FIX / name).read_bytes()


def llm_reply(content, status=200, finish="stop"):
    return httpx.Response(status, json={
        "choices": [{"message": {"role": "assistant", "content": content}, "finish_reason": finish}],
        "usage": {"prompt_tokens": 10, "completion_tokens": 5},
    })


def extractor(handler, **kw) -> OficioExtractor:
    return OficioExtractor(BASE, "qwen-test", transport=httpx.MockTransport(handler), **kw)


def png_bytes(w, h, mode="RGB", color=(255, 255, 255)) -> bytes:
    out = io.BytesIO()
    Image.new(mode, (w, h), color if mode != "L" else 255).save(out, format="PNG")
    return out.getvalue()


def png_header_only(w, h) -> bytes:
    """PNG válido en su cabecera que declara w×h (para bombas de descompresión)."""
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
    ihdr = struct.pack(">IIBBBBB", w, h, 8, 0, 0, 0, 0)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", zlib.compress(b"\x00")) + chunk(b"IEND", b"")


def docx_bytes(build) -> bytes:
    import docx

    d = docx.Document()
    build(d)
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


# --------------------------------------------------------------------------- sniff_mime

@pytest.mark.parametrize("name, mime", [
    ("oficio_texto.pdf", "application/pdf"),
    ("oficio_escaneado.pdf", "application/pdf"),
    ("oficio.png", "image/png"),
    ("oficio_foto.jpg", "image/jpeg"),
    ("oficio_tabla.docx", DOCX_MIME),
])
def test_sniff_real_files(name, mime):
    assert sniff_mime(fixture(name), name) == mime
    assert mime in E.SUPPORTED


def test_sniff_ignores_filename_and_rejects_fakes():
    assert sniff_mime(fixture("oficio_foto.jpg"), "oficio.pdf") == "image/jpeg"
    assert sniff_mime(b"junk\n" * 10 + fixture("oficio_texto.pdf"), "x") == "application/pdf"
    assert sniff_mime(b"", "x.pdf") is None
    assert sniff_mime(b"Esto no es un PDF", "oficio.pdf") is None
    assert sniff_mime(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 100, "viejo.doc") is None  # Word 97
    assert sniff_mime(b"GIF89a" + b"\0" * 20, "x.gif") is None
    assert sniff_mime(b"PK\x03\x04basura que no es zip", "x.docx") is None
    other_zip = io.BytesIO()
    with zipfile.ZipFile(other_zip, "w") as zf:
        zf.writestr("xl/workbook.xml", "<x/>")
    assert sniff_mime(other_zip.getvalue(), "hoja.docx") is None


# --------------------------------------------------------------------------- DOCX

def test_docx_to_text_with_tables_in_order():
    text = E.docx_to_text(fixture("oficio_tabla.docx"))
    lines = text.splitlines()
    assert lines[0] == "[Encabezado de página]"
    assert "FISCALÍA GENERAL DEL ESTADO DE OAXACA" in lines[1]
    assert lines.index("OFICIO No. FGEO/VRC/1187/2026") < lines.index("[Tabla]") < lines.index("[Fin de tabla]")
    table = lines[lines.index("[Tabla]") + 1:lines.index("[Fin de tabla]")]
    assert table == [
        "No. | DESCRIPCIÓN | MARCA | MODELO | SERIE | INVENTARIO",
        "1 | LAPTOP | LENOVO | THINKPAD E14 GEN 2 | PF2ABC12 | FGEO-LAP-000781",
        "2 | NO BREAK | CDP | R-UPR1008 | 2104A0335117 | ",  # celda vacía visible
        "3 | MONITOR | HP | V19 | 1CR0420XYZ | FGEO-MON-002210",
        "4 | TECLADO | HP | KU-1156 | BDMHE0B7Y8K9QZ | S/N",  # combinada verticalmente
    ]
    # el salto de línea de la firma separa nombre y cargo
    assert "LIC. SOFÍA GUADALUPE HERNÁNDEZ VÁSQUEZ\nVICEFISCAL REGIONAL DE LA COSTA" in text
    assert text.index("[Fin de tabla]") < text.index("A T E N T A M E N T E")


def test_docx_horizontal_merge_textbox_and_content_control():
    from docx.oxml import parse_xml

    def build(d):
        d.add_paragraph("Antes de la tabla")
        t = d.add_table(rows=2, cols=3)
        a = t.cell(0, 0).merge(t.cell(0, 1))
        a.text = "TÍTULO COMBINADO"
        t.cell(0, 2).text = "C"
        for j, v in enumerate(["x", "", "z"]):
            t.cell(1, j).text = v
        ns = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
              'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
              'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
              'xmlns:v="urn:schemas-microsoft-com:vml"')
        d.element.body.insert(0, parse_xml(
            f'<w:sdt {ns}><w:sdtContent><w:p><w:r><w:t>Dentro de control</w:t></w:r></w:p>'
            "</w:sdtContent></w:sdt>"))
        p = d.add_paragraph("Párrafo con cuadro:")
        p._p.append(parse_xml(
            f"<w:r {ns}><mc:AlternateContent><mc:Choice Requires=\"wps\"><w:drawing><wps:txbx>"
            "<w:txbxContent><w:p><w:r><w:t>OFICIO NÚM. FGEO/X/1/2026</w:t></w:r></w:p></w:txbxContent>"
            "</wps:txbx></w:drawing></mc:Choice><mc:Fallback><w:pict><v:textbox><w:txbxContent><w:p>"
            "<w:r><w:t>OFICIO NÚM. FGEO/X/1/2026</w:t></w:r></w:p></w:txbxContent></v:textbox></w:pict>"
            "</mc:Fallback></mc:AlternateContent></w:r>"))

    text = E.docx_to_text(docx_bytes(build))
    lines = text.splitlines()
    assert lines[0] == "Dentro de control"
    assert "TÍTULO COMBINADO | C" in lines and "x |  | z" in lines
    assert text.count("OFICIO NÚM. FGEO/X/1/2026") == 1  # el Fallback no se duplica
    assert lines.index("Párrafo con cuadro:") + 1 == lines.index("OFICIO NÚM. FGEO/X/1/2026")


def test_prepare_docx_text_mode():
    prepared = E.prepare(fixture("oficio_tabla.docx"), DOCX_MIME)
    assert prepared.mode == "texto" and prepared.pages >= 1
    assert prepared.text.startswith("CONTENIDO DEL OFICIO")


def test_docx_with_only_a_scanned_image_uses_image_mode(tmp_path):
    scan = tmp_path / "scan.png"
    Image.new("L", (1200, 1550), 255).save(scan)
    content = docx_bytes(lambda d: d.add_picture(str(scan)))
    prepared = E.prepare(content, DOCX_MIME, max_pages=4)
    assert prepared.mode == "imagen" and len(prepared.images) == 1
    with pytest.raises(ExtractionError, match="no tiene texto"):
        E.prepare(docx_bytes(lambda d: d.add_paragraph("hola")), DOCX_MIME)


def test_docx_zip_bomb_guard(monkeypatch):
    monkeypatch.setattr(E, "MAX_ZIP_UNCOMPRESSED", 1000)
    with pytest.raises(ExtractionError, match="demasiado grande"):
        E.prepare(fixture("oficio_tabla.docx"), DOCX_MIME)


def test_corrupt_docx():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("word/document.xml", "<no es xml")
    assert sniff_mime(buf.getvalue()) == DOCX_MIME
    with pytest.raises(ExtractionError, match="Word"):
        E.prepare(buf.getvalue(), DOCX_MIME)


# --------------------------------------------------------------------------- PDF

def test_pdf_with_text_uses_text_mode_and_table_rows():
    prepared = E.prepare(fixture("oficio_texto.pdf"), "application/pdf")
    assert (prepared.mode, prepared.pages, prepared.total_pages) == ("texto", 1, 1)
    assert "OFICIO NÚM.: FGEO/DA/0532/2026" in prepared.text
    assert "1 | COMPUTADORA DE ESCRITORIO | HP | PRODESK 400 G7 | MXL1234ABC | FGEO-CPU-004512" in prepared.text
    assert "3 | MONITOR | DELL | P2219H | CN0KW3G1QDC009A1B2LL | " in prepared.text


def test_pdf_without_text_is_rendered_to_images():
    prepared = E.prepare(fixture("oficio_escaneado.pdf"), "application/pdf")
    assert prepared.mode == "imagen" and prepared.pages == 1
    mime, data = prepared.images[0]
    img = Image.open(io.BytesIO(data))
    assert mime == "image/png" and img.format == "PNG"
    assert 1500 <= max(img.size) <= E.MAX_IMAGE_SIDE  # ~150 dpi en carta (11 in)


def test_scanned_pdf_respects_max_pages():
    doc = pymupdf.open()
    for _ in range(6):
        doc.new_page(width=612, height=792)
    prepared = E.prepare(doc.tobytes(), "application/pdf", max_pages=4)
    assert (prepared.mode, prepared.pages, prepared.total_pages) == ("imagen", 4, 6)
    parts = prepared.content_parts()
    assert [p["type"] for p in parts].count("image_url") == 4
    assert parts[0] == {"type": "text", "text": "Página 1 de 4:"}


def test_short_last_page_keeps_text_mode():
    doc = pymupdf.open(stream=fixture("oficio_texto.pdf"), filetype="pdf")
    doc.new_page(width=612, height=792).insert_text((72, 72), "ATENTAMENTE  LIC. ANA")
    prepared = E.prepare(doc.tobytes(), "application/pdf")
    assert (prepared.mode, prepared.pages) == ("texto", 2)
    assert "--- Página 2 ---" in prepared.text and "LIC. ANA" in prepared.text


def test_pdf_page_covered_by_image_counts_as_scanned():
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    page.insert_image(page.rect, stream=png_bytes(400, 520))
    page.insert_text((72, 72), "Firma electrónica: " + "ABCDEF0123456789" * 20, fontsize=6)
    assert E.prepare(doc.tobytes(), "application/pdf").mode == "imagen"


def test_huge_pdf_page_is_capped():
    doc = pymupdf.open()
    doc.new_page(width=72 * 100, height=72 * 60)  # 100 × 60 pulgadas
    _, data = E.prepare(doc.tobytes(), "application/pdf").images[0]
    assert max(Image.open(io.BytesIO(data)).size) <= E.MAX_IMAGE_SIDE


def test_encrypted_and_corrupt_pdf():
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "secreto " * 50)
    locked = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="x", owner_pw="y")
    with pytest.raises(ExtractionError, match="contraseña"):
        E.prepare(locked, "application/pdf")
    with pytest.raises(ExtractionError, match="PDF"):
        E.prepare(b"%PDF-1.7\nbasura sin estructura", "application/pdf")


# --------------------------------------------------------------------------- imágenes

def test_small_image_is_sent_unchanged():
    content = fixture("oficio_foto.jpg")
    assert E.prepare_image(content) == ("image/jpeg", content)


def test_large_image_is_downscaled():
    mime, data = E.prepare_image(png_bytes(3000, 1500))
    assert mime == "image/png" and Image.open(io.BytesIO(data)).size == (2000, 1000)
    big_jpg = io.BytesIO()
    Image.new("RGB", (4032, 3024), "white").save(big_jpg, format="JPEG")
    mime, data = E.prepare_image(big_jpg.getvalue())
    assert mime == "image/jpeg" and max(Image.open(io.BytesIO(data)).size) == 2000


def test_exif_rotation_and_transparency():
    img = Image.new("RGB", (300, 100), "white")
    exif = img.getexif()
    exif[0x0112] = 6  # girar 90°
    out = io.BytesIO()
    img.save(out, format="JPEG", exif=exif)
    _, data = E.prepare_image(out.getvalue())
    assert Image.open(io.BytesIO(data)).size == (100, 300)
    _, data = E.prepare_image(png_bytes(50, 40, mode="RGBA", color=(0, 0, 0, 0)))
    flat = Image.open(io.BytesIO(data))
    assert flat.mode == "RGB" and flat.getpixel((0, 0)) == (255, 255, 255)


@pytest.mark.parametrize("w, h", [(10_000, 10_000), (60_000, 60_000)])
def test_decompression_bomb_is_rejected_before_decoding(w, h):
    with pytest.raises(ExtractionError, match="demasiado grande"):
        E.prepare_image(png_header_only(w, h))


def test_corrupt_image():
    with pytest.raises(ExtractionError, match="imagen"):
        E.prepare_image(b"\x89PNG\r\n\x1a\nbasura")


# --------------------------------------------------------------------------- llamada al LLM

GOOD = {
    "oficio_recibido": "FGEO/DA/0532/2026", "fecha_oficio_recibido": "2026-08-20",
    "cargo_remitente": "DIRECTORA ADMINISTRATIVA", "nombre_remitente": "MTRA. LAURA ISABEL MENDOZA CRUZ",
    "equipos": [{"tipo_equipo": "MONITOR", "marca": "DELL", "modelo": "P2219H",
                 "numero_serie": "CN0KW3G1QDC009A1B2LL", "numero_inventario": None}],
}


async def test_extract_text_mode_sends_rules_and_schema():
    seen = {}

    def handler(request: httpx.Request):
        seen["url"] = str(request.url)
        seen["auth"] = request.headers.get("authorization")
        seen["body"] = json.loads(request.content)
        return llm_reply(json.dumps(GOOD))

    res = await extractor(handler, timeout_s=30).extract(fixture("oficio_texto.pdf"), "oficio.pdf")
    body = seen["body"]
    assert seen["url"] == f"{BASE}/chat/completions" and seen["auth"] == "Bearer ollama"
    assert body["model"] == "qwen-test" and body["temperature"] == 0 and body["stream"] is False
    assert body["reasoning_effort"] == "none"
    fmt = body["response_format"]
    assert fmt["type"] == "json_schema" and fmt["json_schema"]["strict"] is True
    schema = fmt["json_schema"]["schema"]
    assert schema["properties"]["oficio_recibido"] == {"type": ["string", "null"]}
    assert schema["properties"]["equipos"]["items"]["additionalProperties"] is False
    assert set(schema["required"]) == {"oficio_recibido", "fecha_oficio_recibido", "cargo_remitente",
                                       "nombre_remitente", "equipos"}
    parts = body["messages"][0]["content"]
    assert [p["type"] for p in parts] == ["text", "text"]
    assert "FGEO-CPU-004512" in parts[0]["text"]
    prompt = parts[-1]["text"]
    for rule in ("mes en letra", "FIRMA", "carácter por carácter", "celda está vacía", "No inventes",
                 "2025-08-20"):
        assert rule in prompt
    assert res["mode"] == "texto" and res["pages"] == 1 and isinstance(res["elapsed_s"], float)
    assert res["data"]["fecha_oficio_recibido"] == "20 DE AGOSTO DEL 2026"  # ISO normalizada
    assert res["data"]["equipos"][0]["motivo"] is None


async def test_extract_image_mode_sends_data_url():
    seen = {}

    def handler(request):
        seen["body"] = json.loads(request.content)
        return llm_reply(json.dumps(GOOD))

    res = await extractor(handler).extract(fixture("oficio.png"), "foto.png")
    parts = seen["body"]["messages"][0]["content"]
    assert parts[0]["type"] == "image_url"
    url = parts[0]["image_url"]["url"]
    assert url.startswith("data:image/png;base64,")
    assert base64.b64decode(url.split(",", 1)[1]) == fixture("oficio.png")
    assert res["mode"] == "imagen" and res["pages"] == 1


async def test_extract_json_in_fences():
    reply = "Claro, aquí está:\n```json\n" + json.dumps(GOOD) + "\n```\nSaludos"
    res = await extractor(lambda r: llm_reply(reply)).extract(fixture("oficio.png"), "x.png")
    assert res["data"]["oficio_recibido"] == "FGEO/DA/0532/2026"


async def test_extract_equipos_null():
    res = await extractor(lambda r: llm_reply(json.dumps({**GOOD, "equipos": None}))).extract(
        fixture("oficio.png"), "x.png")
    assert res["data"]["equipos"] == []


@pytest.mark.parametrize("reply, match", [
    (llm_reply("no sé leer esto"), "legibles"),
    (llm_reply(""), "no devolvió datos"),
    (llm_reply("[1, 2, 3]"), "legibles"),
    (llm_reply('{"oficio_recibido": "X", "equipos": [', finish="length"), "incompleta"),
    (llm_reply("ups", status=500), "HTTP 500"),
    (httpx.Response(200, text="<html>proxy</html>"), "legibles"),
    (httpx.Response(200, json={"choices": []}), "legibles"),
])
async def test_extract_bad_responses(reply, match):
    with pytest.raises(ExtractionError, match=match):
        await extractor(lambda r: reply).extract(fixture("oficio.png"), "x.png")


@pytest.mark.parametrize("exc, match", [
    (httpx.ReadTimeout("lento"), "tardó más de 240 s"),
    (httpx.ConnectError("caído"), "conectarme"),
])
async def test_extract_network_errors(exc, match):
    def handler(request):
        raise exc

    with pytest.raises(ExtractionError, match=match):
        await extractor(handler).extract(fixture("oficio.png"), "x.png")


async def test_extract_unsupported_file_does_not_call_llm():
    def boom(request):
        raise AssertionError("no debe llamarse")

    with pytest.raises(ExtractionError, match="Formato no reconocido"):
        await extractor(boom).extract(b"hola", "oficio.pdf")


async def test_llm_calls_are_sequential():
    active, peak = 0, 0

    async def handler(request):
        nonlocal active, peak
        active += 1
        peak = max(peak, active)
        await asyncio.sleep(0.05)
        active -= 1
        return llm_reply(json.dumps(GOOD))

    ex = extractor(handler)
    results = await asyncio.gather(*[ex.extract(fixture("oficio.png"), "x.png") for _ in range(3)])
    assert peak == 1 and len(results) == 3


def test_parse_json_response_variants():
    assert E.parse_json_response('  {"a": 1} ') == {"a": 1}
    assert E.parse_json_response('texto {"a": {"b": 2}} más') == {"a": {"b": 2}}
    assert E.parse_json_response([{"type": "text", "text": '{"a": 1}'}]) == {"a": 1}
    with pytest.raises(ExtractionError):
        E.parse_json_response(None)


def test_ground_truth_fixture_is_consistent():
    # los oficios sintéticos de las pruebas y del experimento comparten la verdad conocida
    assert set(TRUTH) == {"A", "B", "C"}
    text = E.prepare(fixture("oficio_texto.pdf"), "application/pdf").text
    for eq in TRUTH["A"]["equipos"]:
        assert eq["numero_serie"] in text
