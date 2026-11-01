"""Generación del dictamen (app/dictamen/generate.py) y conversión de la plantilla.

Equivalencia con el original: ``tests/fixtures/dictamen/expected_*.json`` guardan la entrada y la
instantánea (texto visible por párrafo y celda en orden, firma de formato por párrafo, tablas y
filas, saltos de página, encabezado y pie) del DOCX que generó el ``generate.js`` ORIGINAL de
dictamen-bot (docxtemplater 3.71.1, plantilla ``config/dictamen/original``, fecha fija con
TZ=America/Mexico_City). Aquí se renderiza lo mismo con docxtpl y se compara con
:func:`docx_snapshot`. Si se cambia la plantilla original hay que regenerarlos con el harness de
Node (``DICTAMEN_TEMPLATE_PATH=… node render_cases.mjs``) y esta misma función.
"""

from __future__ import annotations

import hashlib
import importlib.util
import io
import json
import re
import sys
import unicodedata
import zipfile
from datetime import date, datetime, timezone
from pathlib import Path

import pytest
from docx import Document
from lxml import etree

from app.dictamen import generate
from app.dictamen.generate import (
    DictamenError,
    PlantillaError,
    build_context,
    fecha_larga,
    render_dictamen,
)

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "config" / "dictamen" / "dictamen.docx"
ORIGINAL = ROOT / "config" / "dictamen" / "original" / "dictamen_docxtemplater.docx"
FIXTURES = ROOT / "tests" / "fixtures" / "dictamen"
EXPECTED = sorted(FIXTURES.glob("expected_*.json"))
HOY = date(2026, 10, 5)

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _load_converter():
    spec = importlib.util.spec_from_file_location("convert_dictamen_template", ROOT / "scripts" / "convert_dictamen_template.py")
    if spec.name in sys.modules:
        return sys.modules[spec.name]
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module  # dataclasses busca el módulo en sys.modules
    spec.loader.exec_module(module)
    return module


converter = _load_converter()


# ---------------------------------------------------------------------------
# Instantánea comparable de un DOCX


def _c14n(el) -> bytes:
    return b"" if el is None else etree.tostring(el, method="c14n", exclusive=True, with_comments=False)


def _norm(s: str) -> str:
    return unicodedata.normalize("NFC", s).replace("\xa0", " ")


def _para(p) -> tuple[str, str]:
    """(texto visible, firma de formato = pPr + secuencia (rPr, trozo) fusionando rPr iguales).

    Los saltos de línea son neutrales al formato: docxtemplater los pone en un run sin rPr y docxtpl
    dentro del run del texto; un salto no tiene formato visible propio.
    """
    chunks: list[list] = []
    pending = ""
    for r in p.iter(W + "r"):
        rpr = _c14n(r.find(W + "rPr"))
        for el in r:
            if el.tag in (W + "br", W + "cr"):
                s = "\f" if el.get(W + "type") == "page" else "\n"
                if chunks:
                    chunks[-1][1] += s
                else:
                    pending += s
                continue
            if el.tag == W + "t":
                s = el.text or ""
            elif el.tag == W + "tab":
                s = "\t"
            else:
                continue
            if not s:
                continue
            if chunks and chunks[-1][0] == rpr:
                chunks[-1][1] += s
            else:
                chunks.append([rpr, pending + s])
                pending = ""
    if pending:
        chunks.append([b"", pending])
    text = _norm("".join(c[1] for c in chunks))
    h = hashlib.sha1(_c14n(p.find(W + "pPr")))
    for rpr, s in chunks:
        h.update(b"\0" + rpr + b"\1" + _norm(s).encode())
    return text, h.hexdigest()[:12]


def _blocks(container) -> list[dict]:
    out = []
    for ch in container:
        if ch.tag == W + "p":
            t, f = _para(ch)
            out.append({"p": t, "fmt": f})
        elif ch.tag == W + "tbl":
            rows = []
            for tr in ch.findall(W + "tr"):
                rows.append([[{"p": t, "fmt": f} for t, f in map(_para, tc.findall(W + "p"))] for tc in tr.findall(W + "tc")])
            out.append({"tbl": rows})
        elif ch.tag == W + "sdt":
            out.extend(_blocks(ch.find(W + "sdtContent")))
    return out


def docx_snapshot(content: bytes) -> dict:
    z = zipfile.ZipFile(io.BytesIO(content))
    body = etree.fromstring(z.read("word/document.xml")).find(W + "body")
    blocks = _blocks(body)
    rels = etree.fromstring(z.read("word/_rels/document.xml.rels"))
    targets = {r.get("Id"): "word/" + r.get("Target") for r in rels}
    snap = {
        "blocks": blocks,
        "tables": [[len(row) for row in b["tbl"]] for b in blocks if "tbl" in b],
        "page_breaks": len(body.findall(".//" + W + "pageBreakBefore"))
        + sum(1 for b in body.iter(W + "br") if b.get(W + "type") == "page"),
        "sections": len(body.findall(".//" + W + "sectPr")),
    }
    sect = body.find(W + "sectPr")
    for ref in [*sect.findall(W + "headerReference"), *sect.findall(W + "footerReference")]:
        kind = "header" if ref.tag.endswith("headerReference") else "footer"
        part = etree.fromstring(z.read(targets[ref.get(f"{{{R_NS}}}id")]))
        snap[f"{kind}_{ref.get(W + 'type')}"] = [_para(p)[0] for p in part.iter(W + "p")]
    return snap


def _texts(snap: dict) -> list[str]:
    out = []
    for b in snap["blocks"]:
        if "p" in b:
            out.append(b["p"])
        else:
            out += [" | ".join(" / ".join(c["p"] for c in cell) for cell in row) for row in b["tbl"]]
    return out


def _full_text(content: bytes) -> str:
    return "\n".join(_texts(docx_snapshot(content)))


# ---------------------------------------------------------------------------
# Datos de ejemplo


def _data(n: int = 1, **extra) -> dict:
    equipos = [
        {
            "tipo_equipo": "cpu",
            "marca": "hp",
            "modelo": f"prodesk {i}",
            "numero_serie": f"mxl{i}",
            "numero_inventario": f"fgeo-{i}",
            "motivo": None,
        }
        for i in range(1, n + 1)
    ]
    return {
        "oficio_recibido": "FGEO/DA/0457/2026",
        "fecha_oficio_recibido": "20 de octubre del 2026",
        "cargo_remitente": "Directora de Administración",
        "nombre_remitente": "Mtra. Arima Ojeda Santiago",
        "equipos": equipos,
        **extra,
    }


def _own(**extra) -> dict:
    return {
        "numero_oficio": "fgeo/cgsi/1429/2026",
        "dictamen_inicial": "191",
        "iniciales": "igf/dsa",
        "motivo": "equipo con falla de encendido",
        **extra,
    }


# ---------------------------------------------------------------------------
# Equivalencia contra docxtemplater


def test_hay_fixtures_de_equivalencia():
    assert {p.name for p in EXPECTED} >= {"expected_uno.json", "expected_tres.json", "expected_especiales.json"}


@pytest.mark.parametrize("fixture", EXPECTED, ids=lambda p: p.stem)
def test_equivalente_a_docxtemplater(fixture):
    caso = json.loads(fixture.read_text(encoding="utf-8"))
    entrada = caso["input"]
    content = render_dictamen(entrada["data"], entrada["own"], TEMPLATE, today=date.fromisoformat(entrada["today"]))
    got, expected = docx_snapshot(content), caso["expected"]
    assert _texts(got) == _texts(expected)  # primero el texto, para un diff legible
    assert got == expected


def test_la_comparacion_detecta_diferencias():
    caso = json.loads((FIXTURES / "expected_uno.json").read_text(encoding="utf-8"))
    entrada = caso["input"]
    otro = render_dictamen(entrada["data"], {**entrada["own"], "dictamen_inicial": "192"}, TEMPLATE, today=HOY)
    assert docx_snapshot(otro) != caso["expected"]


# ---------------------------------------------------------------------------
# build_context


def test_build_context_igual_que_generate_js():
    data = _data(3)
    data["equipos"][0]["motivo"] = "tarjeta madre dañada."
    data["equipos"][1]["numero_inventario"] = None
    data["equipos"][2]["numero_inventario"] = "   "
    ctx = build_context(data, _own(motivo="no enciende..."), HOY)
    assert ctx == {
        "numero_oficio": "FGEO/CGSI/1429/2026",
        "fecha_larga": "05 DE OCTUBRE DEL 2026",
        "nombre_remitente": "MTRA. ARIMA OJEDA SANTIAGO",
        "cargo_remitente": "DIRECTORA DE ADMINISTRACIÓN",
        "oficio_recibido": "FGEO/DA/0457/2026",
        "fecha_oficio_recibido": "20 DE OCTUBRE DEL 2026",
        "frase_equipos": "de los siguientes equipos",
        "iniciales": "igf/dsa",
        "otros_componentes": "---",
        "equipos": [
            {"num_dictamen": 191, "tipo_equipo": "CPU", "marca": "HP", "modelo": "PRODESK 1",
             "numero_serie": "MXL1", "numero_inventario": "FGEO-1", "motivo": "TARJETA MADRE DAÑADA"},
            {"num_dictamen": 192, "tipo_equipo": "CPU", "marca": "HP", "modelo": "PRODESK 2",
             "numero_serie": "MXL2", "numero_inventario": "S/N", "motivo": "NO ENCIENDE"},
            {"num_dictamen": 193, "tipo_equipo": "CPU", "marca": "HP", "modelo": "PRODESK 3",
             "numero_serie": "MXL3", "numero_inventario": "S/N", "motivo": "NO ENCIENDE"},
        ],
    }  # fmt: skip


def test_build_context_singular_y_ceros_a_la_izquierda():
    ctx = build_context(_data(1), _own(dictamen_inicial="0099"), HOY)
    assert ctx["frase_equipos"] == "del siguiente equipo"
    assert ctx["equipos"][0]["num_dictamen"] == "0099"
    ctx = build_context(_data(2), _own(dictamen_inicial=" 0099 "), HOY)
    assert [e["num_dictamen"] for e in ctx["equipos"]] == ["0099", "0100"]
    ctx = build_context(_data(2), _own(dictamen_inicial=999), HOY)
    assert [e["num_dictamen"] for e in ctx["equipos"]] == [999, 1000]


def test_build_context_motivo_global_opcional_si_cada_equipo_tiene_el_suyo():
    data = _data(2)
    data["equipos"][0]["motivo"] = "pantalla rota"
    data["equipos"][1]["motivo"] = "no enciende"
    ctx = build_context(data, _own(motivo=None), HOY)
    assert [e["motivo"] for e in ctx["equipos"]] == ["PANTALLA ROTA", "NO ENCIENDE"]


def test_build_context_limpia_controles_y_saltos():
    data = _data(1, cargo_remitente="Directora\r\nde Administración\x07", nombre_remitente="  Ana\x0bRuiz  ")
    data["equipos"][0]["numero_serie"] = "S1\x00\x1f￾"
    ctx = build_context(data, _own(), HOY)
    assert ctx["cargo_remitente"] == "DIRECTORA\nDE ADMINISTRACIÓN"
    assert ctx["nombre_remitente"] == "ANA\nRUIZ"
    assert ctx["equipos"][0]["numero_serie"] == "S1"


def test_build_context_iniciales_tal_cual_y_hoy_por_defecto(monkeypatch):
    monkeypatch.setattr(generate, "hoy_mexico", lambda: date(2027, 1, 2))
    ctx = build_context(_data(1), _own(iniciales="IGF/dsa"), None)
    assert ctx["iniciales"] == "IGF/dsa"
    assert ctx["fecha_larga"] == "02 DE ENERO DEL 2027"


# ---------------------------------------------------------------------------
# fecha_larga


@pytest.mark.parametrize(
    ("d", "esperado"),
    [
        (date(2026, 10, 5), "05 de octubre del 2026"),
        (date(2026, 1, 31), "31 de enero del 2026"),
        (date(2026, 9, 1), "01 de septiembre del 2026"),
        (date(2026, 12, 31), "31 de diciembre del 2026"),
        # 2027-01-01 03:00 UTC sigue siendo 31 de diciembre en Oaxaca (UTC-6)
        (datetime(2027, 1, 1, 3, 0, tzinfo=timezone.utc), "31 de diciembre del 2026"),
        (datetime(2026, 8, 20, 23, 59), "20 de agosto del 2026"),
    ],
)
def test_fecha_larga(d, esperado):
    assert fecha_larga(d) == esperado


def test_fecha_larga_todos_los_meses():
    meses = [fecha_larga(date(2026, m, 1)).split(" de ")[1].split(" del")[0] for m in range(1, 13)]
    assert meses == generate.MESES


def test_hoy_mexico_usa_la_zona_de_mexico(monkeypatch):
    class Reloj(datetime):
        @classmethod
        def now(cls, tz=None):
            return datetime(2027, 1, 1, 4, 30, tzinfo=timezone.utc).astimezone(tz)

    monkeypatch.setattr(generate, "datetime", Reloj)
    assert generate.hoy_mexico() == date(2026, 12, 31)


def test_zona_mexico_sin_tzdata(monkeypatch):
    import zoneinfo

    def falla(_):
        raise zoneinfo.ZoneInfoNotFoundError("sin datos")

    monkeypatch.setattr(zoneinfo, "ZoneInfo", falla)
    zona = generate.zona_mexico()
    assert datetime(2027, 1, 1, 3, 0, tzinfo=timezone.utc).astimezone(zona).date() == date(2026, 12, 31)


# ---------------------------------------------------------------------------
# Errores


def test_sin_equipos_es_error():
    with pytest.raises(DictamenError) as exc:
        build_context(_data(0), _own(), HOY)
    assert exc.value.faltantes == ["equipos"]
    assert "al menos un equipo" in str(exc.value)
    with pytest.raises(DictamenError):
        build_context({k: v for k, v in _data(1).items() if k != "equipos"}, _own(), HOY)


@pytest.mark.parametrize("valor", ["abc", "191A", "191/2026", "19 1", "-5", "1.5", "١٩١", "²"])
def test_dictamen_inicial_solo_digitos(valor):
    with pytest.raises(DictamenError) as exc:
        build_context(_data(1), _own(dictamen_inicial=valor), HOY)
    assert exc.value.faltantes == ["dictamen_inicial"]
    assert "solo dígitos" in str(exc.value)


def test_faltantes_con_nombres_del_bot():
    data = _data(2, oficio_recibido="  ", nombre_remitente=None)
    data["equipos"][1]["marca"] = ""
    data["equipos"][1]["numero_serie"] = None
    with pytest.raises(DictamenError) as exc:
        build_context(data, _own(numero_oficio="", iniciales=None, motivo="."), HOY)
    assert exc.value.faltantes == [
        "oficio_recibido",
        "nombre_remitente",
        "equipo_2_marca",
        "equipo_2_numero_serie",
        "numero_oficio",
        "iniciales",
        "motivo",
    ]
    mensaje = str(exc.value)
    assert mensaje.startswith("Faltan datos para generar el dictamen:")
    assert "marca del equipo 2" in mensaje and "número del oficio recibido" in mensaje


def test_error_combinado_dictamen_inicial_y_faltantes():
    with pytest.raises(DictamenError) as exc:
        build_context(_data(1, cargo_remitente=""), _own(dictamen_inicial="x1"), HOY)
    assert exc.value.faltantes == ["cargo_remitente", "dictamen_inicial"]
    assert "cargo del remitente" in str(exc.value) and "solo dígitos" in str(exc.value)


@pytest.mark.parametrize(
    ("data", "own"),
    [
        ("no es dict", _own()),
        (_data(1), ["no", "es", "dict"]),
        (_data(1, equipos="CPU HP"), _own()),
        (_data(1, equipos=["CPU HP"]), _own()),
        (_data(1, nombre_remitente={"x": 1}), _own()),
        (_data(1, cargo_remitente=True), _own()),
    ],
)
def test_tipos_invalidos(data, own):
    with pytest.raises(DictamenError):
        build_context(data, own, HOY)


def test_dictamen_error_es_value_error():
    assert issubclass(DictamenError, ValueError)


def test_plantilla_inexistente(tmp_path):
    with pytest.raises(PlantillaError):
        render_dictamen(_data(1), _own(), tmp_path / "no-existe.docx", today=HOY)


def test_plantilla_con_error_de_sintaxis(tmp_path):
    doc = Document()
    doc.add_paragraph("{{ numero_oficio ")
    doc.add_paragraph("{% for %}")
    roto = tmp_path / "rota.docx"
    doc.save(roto)
    with pytest.raises(PlantillaError):
        render_dictamen(_data(1), _own(), roto, today=HOY)


def test_los_datos_se_validan_antes_que_la_plantilla(tmp_path):
    with pytest.raises(DictamenError):
        render_dictamen(_data(0), _own(), tmp_path / "no-existe.docx", today=HOY)


# ---------------------------------------------------------------------------
# DOCX generado


@pytest.fixture(scope="module")
def generado() -> bytes:
    data = _data(2, cargo_remitente='Directora "A" & <B>\nSegunda línea')
    data["equipos"][1]["motivo"] = "{{ numero_oficio }} {%p endfor %} {#x#}"
    return render_dictamen(data, _own(), TEMPLATE, today=HOY)


def test_docx_generado_sin_mergefield_mailmerge_ni_metadatos(generado):
    with zipfile.ZipFile(io.BytesIO(generado)) as z:
        assert "word/_rels/settings.xml.rels" not in z.namelist()
        todo = b"".join(z.read(n) for n in z.namelist() if n.endswith((".xml", ".rels")))
    for prohibido in (b"MERGEFIELD", b"w:fldChar", b"w:instrText", b"mailMerge", b"Centinela",
                      b"EFREN", b"SOTO", b"base.dictamen.xlsx", b"python-docx"):  # fmt: skip
        assert prohibido not in todo, prohibido
    props = Document(io.BytesIO(generado)).core_properties
    assert props.author == "" and props.last_modified_by == ""


def test_docx_generado_xml_bien_formado(generado):
    # docxtpl reparsea con recover=True: un XML roto se "repararía" en silencio perdiendo contenido.
    strict = etree.XMLParser(recover=False, huge_tree=True)
    with zipfile.ZipFile(io.BytesIO(generado)) as z:
        assert z.testzip() is None
        for name in z.namelist():
            if name.endswith((".xml", ".rels")):
                etree.fromstring(z.read(name), strict)


def test_docx_generado_abre_con_python_docx_y_sin_etiquetas(generado):
    doc = Document(io.BytesIO(generado))
    assert len(doc.tables) == 2 + 2 * 2  # encabezado, tabla de equipos y 2 tablas por dictamen
    assert len(doc.tables[1].rows) == 1 + 2
    assert doc.tables[1].rows[2].cells[0].text == "192"
    texto = _full_text(generado)
    assert not re.search(r"\{[#/]?[a-z_]+\}", texto)  # nada de docxtemplater sin llenar
    assert 'DIRECTORA "A" & <B>\nSEGUNDA LÍNEA' in texto
    # Los valores nunca se interpretan como plantilla
    assert "dictamen: {{ NUMERO_OFICIO }} {%P ENDFOR %} {#X#}." in texto
    assert texto.count("{{") == 1 and "{%tr" not in texto


def test_render_por_defecto_usa_hoy_en_mexico(monkeypatch):
    monkeypatch.setattr(generate, "hoy_mexico", lambda: date(2026, 2, 3))
    texto = _full_text(render_dictamen(_data(1), _own(), TEMPLATE))
    assert "Oax., a 03 DE FEBRERO DEL 2026" in texto


def test_muchos_equipos():
    content = render_dictamen(_data(30), _own(dictamen_inicial="1"), TEMPLATE, today=HOY)
    snap = docx_snapshot(content)
    assert snap["tables"][1] == [6] * 31
    assert snap["page_breaks"] == 30
    assert _texts(snap).count("MDTG/igf/dsa") == 31
    assert "30 | FGEO-30 | CPU | HP | PRODESK 30 | MXL30" in _texts(snap)


# ---------------------------------------------------------------------------
# Plantilla convertida y script de conversión


def test_plantilla_convertida_al_dia():
    esperado, _ = converter.convert_docx(ORIGINAL.read_bytes())
    assert converter.docx_entries(TEMPLATE.read_bytes()) == converter.docx_entries(esperado), (
        "config/dictamen/dictamen.docx no corresponde al original: corre scripts/convert_dictamen_template.py"
    )


def test_conversion_idempotente():
    original = ORIGINAL.read_bytes()
    una, reporte = converter.convert_docx(original)
    otra, _ = converter.convert_docx(original)
    dos, reporte2 = converter.convert_docx(una)
    assert una == otra and una == dos
    assert reporte.merge_fields == 20 and reporte.mailmerge_removed
    assert reporte.loops == ["equipos (fila de tabla)", "equipos (párrafos)"]
    assert reporte2.merge_fields == 0 and reporte2.tags == 0 and not reporte2.metadata_removed


def test_plantilla_convertida_contenido():
    with zipfile.ZipFile(TEMPLATE) as z:
        doc = z.read("word/document.xml").decode()
        settings = z.read("word/settings.xml").decode()
        core = z.read("docProps/core.xml").decode()
    assert "{%tr for eq in equipos %}" in doc and "{%tr endfor %}" in doc
    assert "{%p for eq in equipos %}" in doc and "{%p endfor %}" in doc
    assert "{{ eq.motivo }}" in doc and "{{ otros_componentes }}" in doc and "{{ fecha_larga }}" in doc
    assert not re.search(r"\{[#/]?[a-z_]+\}", doc)
    assert "fldChar" not in doc and "instrText" not in doc
    assert "mailMerge" not in settings
    assert "creator" not in core and "lastModifiedBy" not in core
    # El texto fijo (con sus erratas) se conserva tal cual.
    assert "le envió un cordial saludo" in doc and "NUMERO DE INVENTARIO" in doc
    assert len(Document(TEMPLATE).tables) == 4  # también abre con python-docx


def _docx_con(*parrafos: list[str]) -> bytes:
    doc = Document()
    for runs in parrafos:
        p = doc.add_paragraph()
        for texto in runs:
            p.add_run(texto)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize(
    ("parrafos", "mensaje"),
    [
        ([["{nom", "bre}"]], "partida"),
        ([["{#equipos}"], ["{#otros}"], ["{/otros}"], ["{/equipos}"]], "anidados"),
        ([["{#equipos}"]], "sin cierre"),
        ([["{a.b}"]], "no soportada"),
        ([["{^equipos}x{/equipos}"]], "no soportada"),
        ([["{#desconocido}"], ["{/desconocido}"]], "sin configuración"),
        ([["{#equipos} y algo"], ["{/equipos}"]], "No sé convertir"),
    ],
)
def test_conversion_rechaza_lo_que_no_sabe_convertir(parrafos, mensaje):
    with pytest.raises(converter.ConversionError, match=mensaje):
        converter.convert_docx(_docx_con(*parrafos))


def test_conversion_bucle_en_linea_y_campos_globales():
    convertido, reporte = converter.convert_docx(
        _docx_con(["Equipos: {#equipos}{marca} ({numero_oficio}), {/equipos}fin"], ["{numero_oficio}"])
    )
    assert reporte.loops == ["equipos (en línea)"]
    doc = zipfile.ZipFile(io.BytesIO(convertido)).read("word/document.xml").decode()
    assert "Equipos: {% for eq in equipos %}{{ eq.marca }} ({{ numero_oficio }}), {% endfor %}fin" in doc
